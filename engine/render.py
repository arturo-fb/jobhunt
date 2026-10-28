"""
Render a cover letter to PDF (reportlab) and DOCX (python-docx),
in a clean one-page business-letter format. No external apps (LibreOffice/Word) needed.
"""
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.enums import TA_LEFT

import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _header_lines(profile):
    links = profile.get("links", {}) or {}
    contact = f'{profile["phone"]} | {profile["email"]}'
    link_bits = []
    if links.get("linkedin"):
        link_bits.append(("LinkedIn", links["linkedin"]))
    if links.get("portfolio"):
        link_bits.append(("Portfolio Website", links["portfolio"]))
    return contact, link_bits


# --------------------------------------------------------------------------- #
def cover_letter_pdf(path, profile, date, recipient, greeting, body, closing="Sincerely,"):
    doc = SimpleDocTemplate(path, pagesize=LETTER,
                            leftMargin=1 * inch, rightMargin=1 * inch,
                            topMargin=1 * inch, bottomMargin=1 * inch)
    styles = getSampleStyleSheet()
    base = ParagraphStyle("body", parent=styles["Normal"], fontName="Helvetica",
                          fontSize=11, leading=15, alignment=TA_LEFT, spaceAfter=11)
    name_st = ParagraphStyle("name", parent=base, fontName="Helvetica-Bold",
                             fontSize=11, spaceAfter=0)
    tight = ParagraphStyle("tight", parent=base, spaceAfter=0)

    contact, link_bits = _header_lines(profile)
    flow = [Paragraph(profile["name"], name_st),
            Paragraph(profile["base_location"], tight),
            Paragraph(contact, tight)]
    if link_bits:
        linked = " | ".join(f'<a href="{u}" color="#0563C1">{t}</a>' for t, u in link_bits)
        flow.append(Paragraph(linked, tight))
    flow.append(Spacer(1, 16))
    flow.append(Paragraph(date, tight))
    flow.append(Spacer(1, 14))
    for line in recipient:
        flow.append(Paragraph(line, tight))
    flow.append(Spacer(1, 14))
    flow.append(Paragraph(greeting, base))
    for para in body:
        flow.append(Paragraph(para, base))
    flow.append(Spacer(1, 6))
    flow.append(Paragraph(closing, tight))
    flow.append(Spacer(1, 18))
    flow.append(Paragraph(profile["name"], tight))
    doc.build(flow)
    return path


# --------------------------------------------------------------------------- #
def cover_letter_docx(path, profile, date, recipient, greeting, body, closing="Sincerely,"):
    d = docx.Document()
    for section in d.sections:
        section.left_margin = section.right_margin = Pt(72)
        section.top_margin = section.bottom_margin = Pt(72)
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def para(text, bold=False, after=6):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        run = p.add_run(text)
        run.bold = bold
        return p

    contact, link_bits = _header_lines(profile)
    para(profile["name"], bold=True, after=0)
    para(profile["base_location"], after=0)
    para(contact, after=0)
    if link_bits:
        para(" | ".join(t for t, _ in link_bits), after=0)
    para("", after=6)
    para(date)
    para("", after=0)
    for line in recipient:
        para(line, after=0)
    para("", after=0)
    para(greeting, after=10)
    for b in body:
        para(b, after=10)
    para(closing, after=0)
    para("", after=12)
    para(profile["name"], after=0)
    d.save(path)
    return path
