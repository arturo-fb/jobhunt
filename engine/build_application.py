#!/usr/bin/env python3
"""
Cover-letter builder — COMPANY-LEVEL, reusable letters.

Cover letters are general about the company (not a single posting), so one
letter is reused across every role you flag at that company. The body is written
by your AI assistant (no paid API); this script handles files + workbook wiring.

Daily flow:
  1. Scan runs (launchd) -> new jobs land with Status "New".
  2. A morning AI scheduled task reviews them and marks the best 5 "Suggested".
  3. You promote the ones you like to "Interested" (and add your own), save.
  4. "make my cover letters" -> one letter per Interested company.

Commands:
  new-today                     list today's New jobs (id, company, title, link)
  suggest <ID> [<ID> ...]       set Status = Suggested for those job ids
  pending                       companies with Interested rows lacking a letter
  context <JOB_ID>              full details for one job
  company-context <COMPANY>     all Interested roles + descriptions for a company
  make-company <body.txt> <COMPANY...>   build ONE company letter, link all its
                                         Interested rows to it
  make-job <JOB_ID> <body.txt>  build a job-SPECIFIC letter (override)
"""
import csv
import datetime as dt
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import openpyxl
from openpyxl.styles import Font
import yaml
import render

XLSX = os.path.join(ROOT, "JobHunt.xlsx")
CONFIG = os.path.join(ROOT, "config", "config.yaml")
CONFIG_EXAMPLE = os.path.join(ROOT, "config", "config.example.yaml")
APPS = os.path.join(ROOT, "Applications")
MASTER = os.path.join(ROOT, "master")
COUNTER = os.path.join(ROOT, "data", "doc_counter.json")
DOCS_INDEX = os.path.join(ROOT, "data", "docs_index.csv")
DESCRIPTIONS = os.path.join(ROOT, "data", "descriptions.json")
LINKF = Font(color="0563C1", underline="single")


def _cfg():
    if not os.path.exists(CONFIG):
        sys.exit("Missing config/config.yaml - copy config/config.example.yaml to "
                 "config/config.yaml and fill in your profile.")
    return yaml.safe_load(open(CONFIG))


def _file_prefix(cfg):
    """'Jane Q. Doe' -> 'JaneQDoe', used for cover-letter file names."""
    return _slug(cfg.get("profile", {}).get("name", ""), 60) or "CoverLetter"


def _load():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Jobs"]
    hdr = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(hdr)}
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = {h: ws.cell(r, idx[h] + 1).value for h in hdr}
        lk = ws.cell(r, idx["Link"] + 1)
        cl = ws.cell(r, idx["Cover Letter"] + 1)
        vals["_link"] = lk.hyperlink.target if lk.hyperlink else ""
        vals["_cl"] = cl.hyperlink.target if cl.hyperlink else ""
        rows.append((r, vals))
    return wb, ws, idx, rows


def _descs():
    return json.load(open(DESCRIPTIONS)) if os.path.exists(DESCRIPTIONS) else {}


def _next_doc_no():
    n = json.load(open(COUNTER)).get("next", 1) if os.path.exists(COUNTER) else 1
    json.dump({"next": n + 1}, open(COUNTER, "w"))
    return n


def _slug(s, n=40):
    return re.sub(r"[^A-Za-z0-9]+", "", s or "")[:n] or "Item"


def _interested(rows, idx):
    return [(r, v) for r, v in rows if (v.get("Status") or "").strip() == "Interested"]


# --------------------------------------------------------------------------- #
def cmd_new_today():
    today = dt.date.today().isoformat()
    _, _, idx, rows = _load()
    new = [(v["Company"], v["Title"], v["Job ID"], v["_link"])
           for _, v in rows
           if (v.get("Status") or "") == "New"
           and str(v.get("Date Found") or "")[:10] == today]
    if not new:
        dates = [str(v.get("Date Found") or "")[:10] for _, v in rows if v.get("Date Found")]
        latest = max(dates) if dates else "none"
        print(f"NO NEW JOBS TODAY ({today}). Latest Date Found in the sheet: {latest}.")
        print("The 6am scan may not have run yet, or there were no new postings.")
        print("STOP: do not analyze or re-suggest old jobs.")
        return
    print(f"{len(new)} New job(s) today ({today}):\n")
    for co, title, jid, link in new:
        print(f"  [{jid}]  {co} — {title}\n        {link}")


def cmd_suggest(ids):
    wb, ws, idx, rows = _load()
    idset = set(ids)
    n = 0
    for r, v in rows:
        if v["Job ID"] in idset:
            ws.cell(r, idx["Status"] + 1, "Suggested")
            n += 1
    wb.save(XLSX)
    print(f"Marked {n} job(s) as Suggested.")


def cmd_pending():
    _, _, idx, rows = _load()
    interested = _interested(rows, idx)
    by_co = {}
    for r, v in interested:
        co = (v["Company"] or "").strip()
        by_co.setdefault(co, {"count": 0, "has_cl": False, "titles": []})
        by_co[co]["count"] += 1
        by_co[co]["titles"].append(v["Title"])
        if v.get("_cl") or v.get("Doc #"):
            by_co[co]["has_cl"] = True
    need = {c: d for c, d in by_co.items() if not d["has_cl"]}
    if not need:
        print("No companies awaiting a cover letter. "
              "(Mark jobs 'Interested' to queue them.)")
        return
    print(f"{len(need)} compan(ies) need a cover letter:\n")
    for co, d in need.items():
        print(f"  {co}  ({d['count']} role(s)):")
        for t in d["titles"]:
            print(f"     - {t}")


def cmd_context(job_id):
    _, _, idx, rows = _load()
    v = next((v for _, v in rows if v["Job ID"] == job_id), None)
    if not v:
        sys.exit(f"Job id not found: {job_id}")
    d = _descs().get(job_id, "")
    print(f"COMPANY:  {v['Company']}\nTITLE:    {v['Title']}\nLOCATION: {v['Location']}"
          f" ({v['Country']})\nLINK:     {v['_link']}\n\nDESCRIPTION:\n{d or '(open the link)'}")


def cmd_company_context(company):
    _, _, idx, rows = _load()
    key = company.strip().lower()
    descs = _descs()
    hits = [v for _, v in rows if (v["Company"] or "").strip().lower() == key
            and (v.get("Status") or "") == "Interested"]
    if not hits:
        hits = [v for _, v in rows if (v["Company"] or "").strip().lower() == key]
    if not hits:
        sys.exit(f"No rows for company: {company}")
    print(f"COMPANY: {hits[0]['Company']}   ({len(hits)} flagged role(s))\n")
    for v in hits:
        print(f"- {v['Title']}  |  {v['Location']}  |  {v['_link']}")
        d = descs.get(v["Job ID"], "")
        if d:
            print(f"    {d[:600]}\n")


def cmd_make_company(body_path, company):
    if not os.path.exists(body_path):
        sys.exit(f"Body file not found: {body_path}")
    cfg = _cfg()
    wb, ws, idx, rows = _load()
    key = company.strip().lower()
    targets = [(r, v) for r, v in rows if (v["Company"] or "").strip().lower() == key
               and (v.get("Status") or "") == "Interested"]
    if not targets:
        targets = [(r, v) for r, v in rows if (v["Company"] or "").strip().lower() == key]
    if not targets:
        sys.exit(f"No rows for company: {company}")
    display = targets[0][1]["Company"].strip()

    body = [p.strip() for p in open(body_path).read().split("\n\n") if p.strip()]
    no = _next_doc_no()
    nnn = f"{no:03d}"
    folder = os.path.join(APPS, _slug(display))
    os.makedirs(folder, exist_ok=True)

    # Resume is NOT copied here — send the single resume you keep in master/.
    date_str = dt.date.today().strftime("%B %d, %Y")
    cl_pdf = os.path.join(folder, f"{_file_prefix(cfg)}.CL.pdf")
    cl_docx = os.path.join(folder, f"{_file_prefix(cfg)}.CL.{nnn}.docx")
    render.cover_letter_pdf(cl_pdf, cfg["profile"], date_str, [display],
                            "Dear Hiring Manager,", body)
    render.cover_letter_docx(cl_docx, cfg["profile"], date_str, [display],
                             "Dear Hiring Manager,", body)

    rel = os.path.relpath(cl_pdf, ROOT)
    for r, v in targets:
        ws.cell(r, idx["Doc #"] + 1, nnn)
        c = ws.cell(r, idx["Cover Letter"] + 1, "CL (company)")
        c.hyperlink = rel
        c.font = LINKF
    wb.save(XLSX)

    _index(nnn, display, f"{len(targets)} role(s)", ",".join(v["Job ID"] for _, v in targets), folder)
    print(f"Company letter built for {display} -> {folder}")
    print(f"  Linked to {len(targets)} Interested role(s). Doc # {nnn}.")


def cmd_make_job(job_id, body_path):
    if not os.path.exists(body_path):
        sys.exit(f"Body file not found: {body_path}")
    cfg = _cfg()
    wb, ws, idx, rows = _load()
    hit = next(((r, v) for r, v in rows if v["Job ID"] == job_id), None)
    if not hit:
        sys.exit(f"Job id not found: {job_id}")
    r, v = hit
    company = (v["Company"] or "").strip()
    title = (v["Title"] or "").strip()
    location = (v["Location"] or "").strip()

    body = [p.strip() for p in open(body_path).read().split("\n\n") if p.strip()]
    no = _next_doc_no()
    nnn = f"{no:03d}"
    folder = os.path.join(APPS, _slug(company), _slug(title, 32))
    os.makedirs(folder, exist_ok=True)

    # Resume is NOT copied here — send the single resume you keep in master/.
    date_str = dt.date.today().strftime("%B %d, %Y")
    recipient = [company] + ([location] if location else [])
    cl_pdf = os.path.join(folder, f"{_file_prefix(cfg)}.CL.pdf")
    cl_docx = os.path.join(folder, f"{_file_prefix(cfg)}.CL.{nnn}.docx")
    render.cover_letter_pdf(cl_pdf, cfg["profile"], date_str, recipient,
                            "Dear Hiring Manager,", body)
    render.cover_letter_docx(cl_docx, cfg["profile"], date_str, recipient,
                             "Dear Hiring Manager,", body)

    rel = os.path.relpath(cl_pdf, ROOT)
    ws.cell(r, idx["Doc #"] + 1, nnn)
    c = ws.cell(r, idx["Cover Letter"] + 1, "CL (job-specific)")
    c.hyperlink = rel
    c.font = LINKF
    wb.save(XLSX)
    _index(nnn, company, title, job_id, folder)
    print(f"Job-specific letter built -> {folder}  (Doc # {nnn})")


def _index(nnn, company, title, job_ids, folder):
    new = not os.path.exists(DOCS_INDEX)
    with open(DOCS_INDEX, "a", newline="") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["Doc #", "Date", "Company", "Title/Roles", "Job IDs", "Folder"])
        w.writerow([nnn, dt.date.today().isoformat(), company, title, job_ids,
                    os.path.relpath(folder, ROOT)])


def main():
    a = sys.argv
    if len(a) < 2:
        print(__doc__); return
    cmd = a[1]
    if cmd == "new-today":
        cmd_new_today()
    elif cmd == "suggest" and len(a) >= 3:
        cmd_suggest(a[2:])
    elif cmd == "pending":
        cmd_pending()
    elif cmd == "context" and len(a) >= 3:
        cmd_context(a[2])
    elif cmd == "company-context" and len(a) >= 3:
        cmd_company_context(" ".join(a[2:]))
    elif cmd == "make-company" and len(a) >= 4:
        cmd_make_company(a[2], " ".join(a[3:]))
    elif cmd == "make-job" and len(a) >= 4:
        cmd_make_job(a[2], a[3])
    else:
        print(__doc__)


if __name__ == "__main__":
    main()
